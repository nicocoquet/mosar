from argparse import ArgumentParser
from pathlib import Path
import sys

from openpyxl import load_workbook

from config import load_config, project_path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from xlsx_preview import prepare_xlsx

CONFIG = load_config()
SHEET = CONFIG["data"]["sheet"]
FORBIDDEN_COLUMNS = set(CONFIG["privacy"]["forbidden_columns"])
ALLOWED_COLUMNS = set(CONFIG["privacy"]["allowed_columns"])
DEFAULT_OUTPUT = project_path(CONFIG["data"]["file"])

if not ALLOWED_COLUMNS:
    raise SystemExit("La liste des colonnes autorisées est vide.")

overlap = FORBIDDEN_COLUMNS & ALLOWED_COLUMNS
if overlap:
    raise SystemExit(
        "Colonnes à la fois autorisées et interdites : "
        + ", ".join(sorted(overlap))
    )

if "ID_Mosar" not in ALLOWED_COLUMNS:
    raise SystemExit("ID_Mosar doit figurer parmi les colonnes autorisées.")


def normalize_header(value):
    return str(value).strip() if value is not None else ""


def get_headers(ws):
    return [normalize_header(cell.value) for cell in ws[1]]


def require_id_column(headers):
    if "ID_Mosar" not in headers:
        raise SystemExit("Colonne obligatoire introuvable : ID_Mosar.")
    if headers.count("ID_Mosar") != 1:
        raise SystemExit("La colonne ID_Mosar apparaît plusieurs fois.")
    return headers.index("ID_Mosar")


def validate_sanitized(path: Path):
    try:
        wb = load_workbook(path, read_only=True, data_only=False)
    except (OSError, ValueError, KeyError) as error:
        raise SystemExit(
            f"Impossible de lire le classeur : {error}"
        ) from error

    try:
        if wb.sheetnames != [SHEET]:
            raise SystemExit(
                f"Le fichier versionné doit contenir uniquement "
                f"l’onglet '{SHEET}'. "
                f"Onglets trouvés : {', '.join(wb.sheetnames)}"
            )

        ws = wb[SHEET]
        headers = get_headers(ws)

        if len(headers) != len(set(headers)):
            raise SystemExit(
                "Contrôle échoué : en-têtes de colonnes dupliqués."
            )

        remaining = sorted(FORBIDDEN_COLUMNS & set(headers))
        if remaining:
            raise SystemExit(
                "Colonnes interdites encore présentes : "
                + ", ".join(remaining)
            )

        unknown = sorted(set(headers) - ALLOWED_COLUMNS)
        if unknown:
            raise SystemExit(
                "Colonnes non autorisées : "
                + ", ".join(unknown)
            )

        id_column = require_id_column(headers)

        for row_number, row in enumerate(
            ws.iter_rows(min_row=2),
            start=2
        ):
            identifier = row[id_column].value

            if identifier is None or not str(identifier).strip():
                raise SystemExit(
                    f"Contrôle échoué : ID_Mosar vide "
                    f"à la ligne {row_number}."
                )

            for cell in row:
                if cell.data_type == "f":
                    raise SystemExit(
                        f"Formule Excel interdite : "
                        f"{SHEET}!{cell.coordinate}"
                    )

    finally:
        wb.close()

    print("Contrôle d’expurgation réussi.")


def sanitize(source: Path, output: Path):
    if source.resolve() == output.resolve():
        raise SystemExit("Le fichier source et le fichier de sortie doivent être différents.")
    try:
        result = prepare_xlsx(
            source.read_bytes(),
            sheet=SHEET,
            allowed_columns=ALLOWED_COLUMNS,
            forbidden_columns=FORBIDDEN_COLUMNS,
        )
    except OSError as error:
        raise SystemExit(f"Impossible de lire le classeur source : {error}") from error
    except ValueError as error:
        raise SystemExit(str(error)) from error

    # Conserver la validation et le remplacement atomique déjà utilisés par la CI.
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.name + ".tmp.xlsx")
    try:
        temporary.write_bytes(result.content)
        validate_sanitized(temporary)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)

    report = result.report
    print(f"Onglets exclus ({len(report['sheets_removed'])}) : {', '.join(report['sheets_removed']) or 'aucun'}")
    print(f"Colonnes confidentielles supprimées ({len(report['private_columns_removed'])}) : {', '.join(report['private_columns_removed']) or 'aucune'}")
    print(f"Colonnes non autorisées supprimées ({len(report['unknown_columns_removed'])}) : {', '.join(report['unknown_columns_removed']) or 'aucune'}")
    print(f"Lignes sans ID_Mosar exclues : {report['rows_without_id_removed']}")
    print(f"Lignes de données conservées : {report['rows_retained']}")
    print(f"Fichier expurgé créé : {output}")


def main():
    parser = ArgumentParser()
    parser.add_argument("source", nargs="?", type=Path)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--check", nargs="?", type=Path, const=DEFAULT_OUTPUT)
    args = parser.parse_args()
    if args.check:
        validate_sanitized(args.check)
        return
    if not args.source:
        parser.error("un fichier source est requis sauf avec --check")
    sanitize(args.source, args.output)


if __name__ == "__main__":
    main()
