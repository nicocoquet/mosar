from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
import math
import re

from openpyxl import load_workbook

from mosar_model import (
    ALLOWED_STRUCTURES,
    DISCIPLINE_MAP,
    FORMAT_MAP,
    LICENCE_MAP,
    PROXIMITY_LEVELS,
    creation_period,
    date_access_group,
    group_structure_for_access,
    normalize_structure,
    periodicity_label,
)


@dataclass(frozen=True)
class StatisticsSource:
    xlsx: Path
    sheet: str
    columns: dict


def norm(v):
    return "" if v is None else re.sub(r"\s+", " ", str(v).strip())


def rpct(v, total):
    return math.floor(v * 100 / total + 0.5) if total else 0


def read_rows(source):

    wb = load_workbook(source.xlsx, read_only=True, data_only=True)
    if source.sheet not in wb.sheetnames:
        raise SystemExit(f"Onglet attendu introuvable : {source.sheet}")
    ws = wb[source.sheet]

    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    required = [
        source.columns["journal_name"],
        source.columns["proximity"],
        source.columns["open_access"],
        source.columns["format"],
        source.columns["creation_year"],
        source.columns["periodicity"],
        source.columns["discipline"],
        *([source.columns["sector"]] if "sector" in source.columns else []),
        source.columns["publisher_type"],
        source.columns["licence"],
    ]
    missing = [x for x in required if x not in headers]
    if missing:
        raise SystemExit("Colonnes manquantes : " + ", ".join(missing))

    rows = []
    for row_number, values in enumerate(
        ws.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        row = dict(zip(headers, values))
        if norm(row[source.columns["journal_name"]]):
            row["_row"] = row_number
            rows.append(row)
    return rows


def access(v, row_number):
    x = norm(v).casefold()
    if x == "oui":
        return "open"
    if x == "non":
        return "restricted"
    raise SystemExit(f"Accès ouvert invalide ligne {row_number}: {v!r}")


def year(v, row_number):
    if v is None or norm(v) in ("", "?"):
        return None

    try:
        y = int(float(v))
    except (TypeError, ValueError):
        raise SystemExit(f"Année de création invalide ligne {row_number}: {v!r}")

    if not 1500 <= y <= 2100:
        raise SystemExit(f"Année de création hors plage ligne {row_number}: {y}")

    return y


def periodicity(v, row_number):
    x = norm(v).casefold()
    try:
        return periodicity_label(x)
    except (TypeError, ValueError):
        raise SystemExit(f"Périodicité invalide ligne {row_number}: {v!r}")



def creation_quantiles(year_access, groups=4):
    """Contiguous quantile classes, keeping all journals of a year together."""
    years = sorted(year_access)
    if not years:
        return []
    weights = [sum(year_access[y].values()) for y in years]
    total = sum(weights)
    # Dynamic programming: minimize squared deviation from target group size.
    k = min(groups, len(years))
    prefix = [0]
    for weight in weights:
        prefix.append(prefix[-1] + weight)
    target = total / k
    dp = [[float("inf")] * (len(years) + 1) for _ in range(k + 1)]
    paths = [[None] * (len(years) + 1) for _ in range(k + 1)]
    dp[0][0] = 0
    for count in range(1, k + 1):
        for end in range(count, len(years) + 1):
            for start in range(count - 1, end):
                score = dp[count - 1][start] + (prefix[end] - prefix[start] - target) ** 2
                if score < dp[count][end]:
                    dp[count][end], paths[count][end] = score, start
    boundaries = []
    end = len(years)
    for count in range(k, 0, -1):
        start = paths[count][end]
        boundaries.append((start, end))
        end = start
    boundaries.reverse()
    result = []
    for start, end in boundaries:
        counts = Counter()
        for y in years[start:end]:
            counts.update(year_access[y])
        result.append((years[start], years[end - 1], dict(counts)))
    assert sum(sum(counts.values()) for _, _, counts in result) == total
    return result


def derive(rows, source):
    columns = source.columns
    total = len(rows)
    levels, formats, years, periods = Counter(), Counter(), Counter(), Counter()
    periodicities, disciplines, structures, rights = Counter(), Counter(), Counter(), Counter()
    sectors = Counter()
    level_access, format_access, date_access, structure_access = defaultdict(Counter), defaultdict(Counter), defaultdict(Counter), defaultdict(Counter)
    year_access = defaultdict(Counter)

    for row in rows:
        rn = row["_row"]
        a = access(row[columns["open_access"]], rn)
        try:
            level = int(row[columns["proximity"]])
        except (TypeError, ValueError):
            raise SystemExit(f"Niveau de rattachement invalide ligne {rn}")
        if level not in PROXIMITY_LEVELS:
            raise SystemExit(f"Niveau de rattachement invalide ligne {rn}: {level}")
        levels[level] += 1
        level_access[level][a] += 1

        f = norm(row[columns["format"]])
        if f not in FORMAT_MAP:
            raise SystemExit(f"Format inattendu ligne {rn}: {f!r}")
        f = FORMAT_MAP[f]
        formats[f] += 1
        format_access[f][a] += 1

        y = year(row[columns["creation_year"]], rn)
        if y is not None:
            years[y] += 1
            year_access[y][a] += 1
            periods[creation_period(y)] += 1
            date_group = date_access_group(y)
            date_access[date_group][a] += 1

        periodicity_value = row[columns["periodicity"]]
        if periodicity_value is not None and norm(periodicity_value) not in ("", "?"):
            periodicities[periodicity(periodicity_value, rn)] += 1

        discipline = norm(row[columns["discipline"]])
        if discipline not in DISCIPLINE_MAP:
            raise SystemExit(f"Discipline non référencée ligne {rn}: {discipline!r}")
        disciplines[DISCIPLINE_MAP[discipline]] += 1

        raw_sector = norm(row[columns["sector"]]) if "sector" in columns else ""
        sector_values = {part.strip().upper() for part in raw_sector.split("|") if part.strip()}
        if sector_values == {"SHS"}:
            sectors["SHS"] += 1
        elif sector_values == {"STM"}:
            sectors["STM"] += 1
        elif sector_values == {"SHS", "STM"}:
            sectors["SHS + STM"] += 1
        elif raw_sector not in ("", "?"):
            raise SystemExit(f"Secteur inattendu ligne {rn}: {raw_sector!r}")

        structure = norm(row[columns["publisher_type"]])
        if structure not in ALLOWED_STRUCTURES:
            raise SystemExit(f"Type d'éditeur inattendu ligne {rn}: {structure!r}")
        structures[normalize_structure(structure)] += 1
        grouped_structure = group_structure_for_access(structure)
        structure_access[grouped_structure][a] += 1

        licence = norm(row[columns["licence"]])
        if licence:
            if licence not in LICENCE_MAP:
                raise SystemExit(f"Licence non référencée ligne {rn}: {licence!r}")
            rights[LICENCE_MAP[licence]] += 1

    complete_counters = [
        ("niveaux", levels),
        ("formats", formats),
        ("disciplines", disciplines),
        ("structures", structures),
    ]

    for name, counter in complete_counters:
        if sum(counter.values()) != total:
            raise SystemExit(
                f"Contrôle de total échoué pour {name}: "
                f"{sum(counter.values())}/{total}"
            )

    partial_counters = [
        ("années", years),
        ("périodicités", periodicities),
        ("droits", rights),
        ("secteurs", sectors),
    ]

    for name, counter in partial_counters:
        if sum(counter.values()) > total:
            raise SystemExit(
                f"Contrôle de total incohérent pour {name}: "
                f"{sum(counter.values())}/{total}"
            )

    return {
        "total": total,
        "levels": levels,
        "level_access": level_access,
        "formats": formats,
        "format_access": format_access,
        "years": years,
        "year_access": year_access,
        "creation_quantiles": creation_quantiles(year_access),
        "periods": periods,
        "date_access": date_access,
        "periodicities": periodicities,
        "disciplines": disciplines,
        "sectors": sectors,
        "structures": structures,
        "struct_access": structure_access,
        "rights": rights,
    }
