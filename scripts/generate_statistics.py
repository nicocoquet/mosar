from collections import Counter
from pathlib import Path
import csv
import math

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "Recensement-revues-stat.xlsx"
SHEET = "recencement"
OUTPUTS = {"fr": ROOT / "docs/statistiques.md", "en": ROOT / "docs/statistiques.en.md"}
DOWNLOADS = ROOT / "docs/downloads"
BLUES = ["#07558c", "#367fb8", "#67a9df", "#a7d5f5"]

DEFINITIONS = {
    "fr": {
        1: "revues publiées par la MSH Mondes, les Presses universitaires de Nanterre et les Éditions de la Sorbonne",
        2: "revues portées ou hébergées par un laboratoire des universités Paris 1 ou Paris Nanterre",
        3: "revues dont le.a rédacteur.rice en chef est membre des universités Paris 1 ou Paris Nanterre, qui comprennent au moins un.e membre du comité de rédaction restreint (≤ 8) de Paris 1 ou Paris Nanterre ou si la revue est soutenue financièrement par une UR",
        4: "revues dont au moins un.e membre appartient à l’université Paris 1 ou Paris Nanterre (comité de rédaction ou comité scientifique)",
    },
    "en": {
        1: "journals published by MSH Mondes, Presses universitaires de Nanterre or Éditions de la Sorbonne",
        2: "journals hosted or supported by a research unit affiliated with Paris 1 or Paris Nanterre University",
        3: "journals whose editor-in-chief is affiliated with Paris 1 or Paris Nanterre University, which include at least one member from Paris 1 or Paris Nanterre in a small editorial board (≤ 8), or which receive financial support from a research unit",
        4: "journals with at least one member affiliated with Paris 1 or Paris Nanterre University on the editorial or scientific board",
    },
}

TITLES = {
    "fr": "Répartition des revues selon le degré de proximité avec les universités Paris Nanterre, Paris 1 et la MSH Mondes",
    "en": "Distribution of journals by degree of proximity to Paris Nanterre University, Paris 1 University and MSH Mondes",
}

CSV_ICON = '''<span class="chart-download-icon" aria-hidden="true"><svg viewBox="0 0 48 56"><path d="M8 2h21l11 11v39a2 2 0 0 1-2 2H8a2 2 0 0 1-2-2V4a2 2 0 0 1 2-2Z"/><path d="M29 2v12h11M14 26h20v17H14zM14 32h20M14 38h20M21 26v17M28 26v17"/></svg></span>'''
JPG_ICON = '''<span class="chart-download-icon" aria-hidden="true"><svg viewBox="0 0 48 56"><path d="M8 2h21l11 11v39a2 2 0 0 1-2 2H8a2 2 0 0 1-2-2V4a2 2 0 0 1 2-2Z"/><path d="M29 2v12h11M14 42l7-8 5 5 5-7 7 10M18 25.5a3 3 0 1 0 0 .1Z"/></svg></span>'''
SVG_ICON = '''<span class="chart-download-icon" aria-hidden="true"><svg viewBox="0 0 48 56"><path d="M8 2h21l11 11v39a2 2 0 0 1-2 2H8a2 2 0 0 1-2-2V4a2 2 0 0 1 2-2Z"/><path d="M29 2v12h11M16 39l6-10 6 10M18.5 35h7M31 28h4v12h-4M31 34h3"/></svg></span>'''


def percent(value, total):
    return value * 100 / total if total else 0


def rounded_percent(value, total):
    return math.floor(percent(value, total) + 0.5)


def polar(cx, cy, r, angle):
    a = math.radians(angle - 90)
    return cx + r * math.cos(a), cy + r * math.sin(a)


def sector(cx, cy, r, start, end):
    x1, y1 = polar(cx, cy, r, start)
    x2, y2 = polar(cx, cy, r, end)
    large = 1 if end - start > 180 else 0
    return f"M {cx},{cy} L {x1:.2f},{y1:.2f} A {r},{r} 0 {large},1 {x2:.2f},{y2:.2f} Z"


def read_counts():
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"Onglet attendu introuvable : {SHEET}")
    ws = wb[SHEET]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    name_col = headers.index("Nom revue")
    level_col = headers.index("Niveau rattachement")
    counts = Counter()
    total = 0
    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row[name_col] or not str(row[name_col]).strip():
            continue
        total += 1
        try:
            level = int(row[level_col])
        except (TypeError, ValueError):
            raise SystemExit(f"Niveau de rattachement invalide ligne {row_number}")
        if level not in (1, 2, 3, 4):
            raise SystemExit(f"Niveau de rattachement invalide ligne {row_number}: {level}")
        counts[level] += 1
    if sum(counts.values()) != total:
        raise SystemExit("Toutes les revues doivent avoir un niveau de rattachement valide.")
    return {n: counts[n] for n in (1, 2, 3, 4)}, total


def fmt_pct(value, total, lang):
    value = rounded_percent(value, total)
    return f"{value} %" if lang == "fr" else f"{value}%"


def chart_svg_markup(counts, total, lang, standalone=False):
    cx, cy, r = 205, 215, 155
    current = 0
    label = "Niveau" if lang == "fr" else "Level"
    parts = []
    for i, level in enumerate((1, 2, 3, 4)):
        value = counts[level]
        end = current + percent(value, total) * 3.6
        parts.append(f'<path d="{sector(cx, cy, r, current, end)}" fill="{BLUES[i]}" stroke="#ffffff" stroke-width="2.2"/>')
        tx, ty = polar(cx, cy, r * .62, current + (end-current)/2)
        color = "#fff" if level in (1, 2) else "#08264a"
        parts.append(
            f'<text x="{tx:.1f}" y="{ty-14:.1f}" text-anchor="middle" fill="{color}" font-family="Arial, sans-serif" font-size="16" font-weight="700">'
            f'<tspan x="{tx:.1f}">{label} {level}</tspan>'
            f'<tspan x="{tx:.1f}" dy="22">({value})</tspan>'
            f'<tspan x="{tx:.1f}" dy="22">{fmt_pct(value,total,lang)}</tspan></text>'
        )
        current = end

    combined = counts[1] + counts[2]
    combined_label = "Niveau 1 et 2" if lang == "fr" else "Levels 1 and 2"
    content = f'''<g>{''.join(parts)}</g>
<path d="M375,105 C395,105 395,125 395,145 C395,170 410,175 420,175 C410,175 395,180 395,205 C395,225 395,245 375,245" fill="none" stroke="#367fb8" stroke-width="4" stroke-linecap="round"/>
<g>
  <rect x="440" y="125" width="155" height="110" rx="10" fill="#f2f7fc" stroke="#dce7f2" stroke-width="1.2"/>
  <text x="517" y="160" text-anchor="middle" fill="#08264a" font-family="Arial, sans-serif" font-size="16" font-weight="700">{combined_label}</text>
  <text x="517" y="195" text-anchor="middle" fill="#08264a" font-family="Arial, sans-serif" font-size="27" font-weight="700">({combined})</text>
  <text x="517" y="224" text-anchor="middle" fill="#08264a" font-family="Arial, sans-serif" font-size="27" font-weight="700">{fmt_pct(combined,total,lang)}</text>
</g>'''
    if standalone:
        return f'''<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 620 430" width="620" height="430" role="img">
{content}
</svg>
'''
    return f'''<div class="chart-shell">
<svg class="proximity-chart" viewBox="0 0 620 430" role="img" aria-label="{TITLES[lang]}">
{content}
</svg>
</div>'''


def legend(lang):
    label = "Niveau" if lang == "fr" else "Level"
    items = "".join(
        f'<div class="legend-item"><span class="legend-dot level-{n}"></span><p><strong>{label} {n} :</strong> {DEFINITIONS[lang][n]}</p></div>'
        for n in (1, 2, 3, 4)
    )
    return f'<section class="level-legend">{items}</section>'


def export_csv(counts, total, lang):
    path = DOWNLOADS / f"proximite-revues.{lang}.csv"
    headers = ["Niveau", "Nombre", "Pourcentage"] if lang == "fr" else ["Level", "Count", "Percentage"]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(headers)
        for level in (1, 2, 3, 4):
            writer.writerow([level, counts[level], rounded_percent(counts[level], total)])
    return path


def export_jpeg(counts, total, lang):
    path = DOWNLOADS / f"proximite-revues.{lang}.jpg"
    labels = [("Niveau" if lang == "fr" else "Level") + f" {n}" for n in (1, 2, 3, 4)]
    values = [counts[n] for n in (1, 2, 3, 4)]

    fig, ax = plt.subplots(figsize=(7.2, 4.8), dpi=170)
    fig.patch.set_facecolor("white")
    ax.set_aspect("equal")

    wedges, _ = ax.pie(values, colors=BLUES, startangle=90, counterclock=False, wedgeprops={"linewidth": 1.2, "edgecolor": "white"})
    for wedge, level, value in zip(wedges, (1, 2, 3, 4), values):
        angle = (wedge.theta1 + wedge.theta2) / 2
        x = .62 * math.cos(math.radians(angle))
        y = .62 * math.sin(math.radians(angle))
        color = "white" if level in (1, 2) else "#08264a"
        ax.text(x, y, f"{labels[level-1]}\n({value})\n{fmt_pct(value,total,lang)}", ha="center", va="center", fontsize=9, fontweight="bold", color=color)

    combined = counts[1] + counts[2]
    combined_label = "Niveau 1 et 2" if lang == "fr" else "Levels 1 and 2"
    ax.text(1.35, .45, f"{combined_label}\n({combined})\n{fmt_pct(combined,total,lang)}", ha="center", va="center", fontsize=10, fontweight="bold", color="#08264a", bbox={"boxstyle": "round,pad=.45", "facecolor": "#f2f7fc", "edgecolor": "#dce7f2"})

    ax.set_xlim(-1.2, 2.05)
    ax.set_ylim(-1.18, 1.18)
    ax.axis("off")
    fig.savefig(path, format="jpg", bbox_inches="tight", pad_inches=.08, facecolor="white")
    plt.close(fig)
    return path


def export_svg(counts, total, lang):
    path = DOWNLOADS / f"proximite-revues.{lang}.svg"
    path.write_text(chart_svg_markup(counts, total, lang, standalone=True), encoding="utf-8")
    return path


def page(counts, total, lang):
    csv_label = "Télécharger les données" if lang == "fr" else "Download data"
    jpg_label = "Télécharger le graphique" if lang == "fr" else "Download chart"
    svg_label = "Télécharger le graphique" if lang == "fr" else "Download chart"
    jpg_title = "Télécharger le graphique au format JPEG" if lang == "fr" else "Download chart as JPEG"
    svg_title = "Télécharger le graphique au format SVG" if lang == "fr" else "Download chart as SVG"
    return f'''<h1 class="statistics-title">{TITLES[lang]}</h1>

<div class="chart-block">
  <div class="chart-actions">
    <a class="chart-download" href="../downloads/proximite-revues.csv" download>{CSV_ICON}<span>{csv_label}</span></a>
    <a class="chart-download" href="../downloads/proximite-revues.jpg" download title="{jpg_title}" aria-label="{jpg_title}">{JPG_ICON}<span>{jpg_label}</span></a>
    <a class="chart-download" href="../downloads/proximite-revues.svg" download title="{svg_title}" aria-label="{svg_title}">{SVG_ICON}<span>{svg_label}</span></a>
  </div>
  <div class="chart-layout">
    {chart_svg_markup(counts, total, lang)}
    {legend(lang)}
  </div>
</div>
'''


def main():
    counts, total = read_counts()
    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    for lang, output in OUTPUTS.items():
        export_csv(counts, total, lang)
        export_jpeg(counts, total, lang)
        export_svg(counts, total, lang)
        output.write_text(page(counts, total, lang), encoding="utf-8")
    print(
        f"Corpus : {total} | N1={counts[1]} N2={counts[2]} N3={counts[3]} N4={counts[4]} | "
        f"N1+N2={counts[1]+counts[2]} ({rounded_percent(counts[1]+counts[2], total)} %)"
    )


if __name__ == "__main__":
    main()
