"""Moteur d'expurgation XLSX en mémoire, sans accès à GitHub ni au disque.

La sortie ne contient que les colonnes autorisées. Le rapport n'expose
jamais les valeurs des colonnes supprimées.
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook


@dataclass(frozen=True)
class Preview:
    content: bytes
    report: dict[str, Any]


def _headers(sheet):
    return [str(c.value).strip() if c.value is not None else "" for c in sheet[1]]


def prepare_xlsx(source: bytes, *, sheet: str, allowed_columns: set[str],
                 forbidden_columns: set[str]) -> Preview:
    """Expurge un XLSX reçu sous forme d'octets et produit un rapport sans valeurs privées."""
    if not source:
        raise ValueError("Classeur vide.")
    if not allowed_columns or "ID_Mosar" not in allowed_columns:
        raise ValueError("Configuration invalide : ID_Mosar doit être autorisé.")
    if allowed_columns & forbidden_columns:
        raise ValueError("Configuration invalide : colonnes autorisées et interdites communes.")
    try:
        book = load_workbook(BytesIO(source), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("Impossible de lire le classeur XLSX.") from exc
    clean = Workbook()
    try:
        if sheet not in book.sheetnames:
            raise ValueError(f"Onglet attendu introuvable : {sheet}")
        ws = book[sheet]
        headers = _headers(ws)
        if len(headers) != len(set(headers)):
            raise ValueError("En-têtes de colonnes dupliqués dans le classeur source.")
        if headers.count("ID_Mosar") != 1:
            raise ValueError("Colonne obligatoire introuvable : ID_Mosar.")
        id_index = headers.index("ID_Mosar")
        indexes = [i for i, name in enumerate(headers) if name in allowed_columns]
        removed_private = sorted(set(headers) & forbidden_columns)
        removed_unknown = sorted(set(headers) - allowed_columns - forbidden_columns)
        excluded_sheets = [name for name in book.sheetnames if name != sheet]
        output_sheet = clean.active
        output_sheet.title = sheet
        output_sheet.append([headers[i] for i in indexes])
        retained = excluded = 0
        for row_number, row in enumerate(ws.iter_rows(min_row=2), start=2):
            identifier = row[id_index].value
            if identifier is None or not str(identifier).strip():
                excluded += 1
                continue
            values = []
            for index in indexes:
                cell = row[index]
                value = cell.value
                if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                    raise ValueError(
                        f"Formule Excel interdite : {sheet}, ligne {row_number}, "
                        f"colonne '{headers[index]}'."
                    )
                values.append(value)
            output_sheet.append(values)
            retained += 1
        if not retained:
            raise ValueError("Aucune revue avec ID_Mosar : classeur refusé.")
        stream = BytesIO()
        clean.save(stream)
        return Preview(
            content=stream.getvalue(),
            report={
                "sheets_removed": excluded_sheets,
                "private_columns_removed": removed_private,
                "unknown_columns_removed": removed_unknown,
                "rows_without_id_removed": excluded,
                "rows_retained": retained,
                "columns_retained": len(indexes),
            },
        )
    finally:
        book.close()
        clean.close()
