#!/usr/bin/env python3
"""Explorateur des revues : métadonnées Mir@bel + enrichissements analytiques Mosar."""
from __future__ import annotations

import csv
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from config import load_config, project_path
from mosar_model import PROXIMITY_LEVELS, proximity_sort_key
from journals_model import build_record

CONFIG = load_config()
COLUMNS = CONFIG["columns"]
GRAPPE_ID = CONFIG["mirabel"]["cluster_id"]
GRAPPE_NAME = CONFIG["mirabel"]["cluster_name"]
GRAPPE_URL = CONFIG["mirabel"]["cluster_url"]
API = CONFIG["mirabel"]["api"].rstrip("/")
XLSX = project_path(CONFIG["data"]["file"])
SHEET = CONFIG["data"]["sheet"]
PROJECT_NAME = CONFIG["project"]["name"]
SITE_URL = CONFIG["project"]["site_url"]
JOURNALS_PUBLICATION = CONFIG["publication"]["journals"]
JOURNALS_PAGE = JOURNALS_PUBLICATION["pages"]["fr"]
FALLBACK_CONFIG = JOURNALS_PUBLICATION["fallback"]
FALLBACK_ENABLED = bool(FALLBACK_CONFIG.get("enabled", False))
FALLBACK = project_path(FALLBACK_CONFIG["file"])
OUT_JSON = project_path(JOURNALS_PUBLICATION["data_file"])
OUT_REPORT = project_path(JOURNALS_PUBLICATION["report_file"])
INTRO_FILE = project_path(JOURNALS_PAGE["intro"])
OUT_MD = project_path(JOURNALS_PAGE["output"])


def api_json(path: str):
    req = Request(f"{API}{path}", headers={"User-Agent": f"{PROJECT_NAME}/1.0 (+{SITE_URL})"})
    with urlopen(req, timeout=45) as response:
        return json.load(response)


def fetch_titles() -> list[dict]:
    records, offset = [], 0
    while True:
        q = urlencode({"grappeid": GRAPPE_ID, "actif": 1, "offset": offset})
        batch = api_json(f"/titres?{q}")
        if not isinstance(batch, list):
            raise RuntimeError("Réponse inattendue de /titres")
        records.extend(batch)
        if len(batch) < 1000:
            return records
        offset += 1000


def load_mosar() -> tuple[dict[int, dict], str]:
    try:
        from openpyxl import load_workbook
        ws = load_workbook(XLSX, data_only=True, read_only=True)[SHEET]
        rows = ws.iter_rows(values_only=True)
        headers = [str(v).strip() if v is not None else "" for v in next(rows)]
        if COLUMNS["mirabel_id"] in headers:
            pos = {h: i for i, h in enumerate(headers)}
            fields = {
                "nom_revue": COLUMNS["journal_name"],
                "niveau_rattachement": COLUMNS["proximity"],
                "acces_ouvert": COLUMNS["open_access"],
                "discipline": COLUMNS["discipline"],
                "periodicite": COLUMNS["periodicity"],
                "format": COLUMNS["format"],
                "type_editeur": COLUMNS["publisher_type"],
                "licence": COLUMNS["licence"],
            }
            out = {}
            for row in rows:
                raw = row[pos[COLUMNS["mirabel_id"]]]
                if raw in (None, ""): continue
                out[int(raw)] = {key: row[pos[col]] if col in pos else None for key, col in fields.items()}
            return out, f"XLSX {PROJECT_NAME}"
    except Exception as exc:
        if not FALLBACK_ENABLED:
            raise RuntimeError(
                "Lecture XLSX impossible pour la jointure Mosar et fallback CSV désactivé"
            ) from exc
        print(
            f"Lecture XLSX impossible pour la jointure ({exc}); "
            f"recours au CSV configuré : {FALLBACK}."
        )

    out = {}
    with FALLBACK.open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            rid = int(row.pop("mirabel_id"))
            out[rid] = row
    return out, f"CSV de fallback ({FALLBACK})"


def load_intro() -> str:
    if not INTRO_FILE.is_file():
        raise SystemExit(
            f"Introduction éditoriale de la page Revues introuvable : {INTRO_FILE}"
        )
    return INTRO_FILE.read_text(encoding="utf-8").strip()


def main():
    titles = fetch_titles()
    themes_payload = api_json(f"/themes/grappe/{GRAPPE_ID}")
    themes_by_revue = {int(x["revueid"]): [t.get("nom") for t in x.get("themes", []) if t.get("nom")] for x in themes_payload}
    mosar, mosar_source = load_mosar()
    by_revue = {}; duplicate_active_titles = Counter()
    for t in titles:
        rid = t.get("revueid")
        if not rid: continue
        rid = int(rid); duplicate_active_titles[rid] += 1
        if rid in by_revue: continue
        by_revue[rid] = build_record(
            t,
            themes_by_revue.get(rid, []),
            mosar.get(rid),
        )
    records = sorted(by_revue.values(), key=proximity_sort_key)
    cluster_ids=set(by_revue); mosar_ids=set(mosar); matched=sorted(cluster_ids & mosar_ids)
    report={"generated_at":datetime.now(timezone.utc).isoformat(timespec="seconds"),"cluster_id":GRAPPE_ID,"mirabel_active_titles":len(titles),"mirabel_unique_reviews":len(records),"mosar_source":mosar_source,"mosar_ids":len(mosar_ids),"matched_ids":matched,"matched_count":len(matched),"mirabel_without_mosar_count":len(cluster_ids-mosar_ids),"mosar_outside_cluster_ids":sorted(mosar_ids-cluster_ids),"duplicate_active_title_review_ids":sorted(rid for rid,n in duplicate_active_titles.items() if n>1)}
    OUT_JSON.parent.mkdir(parents=True,exist_ok=True); OUT_REPORT.parent.mkdir(parents=True,exist_ok=True)
    OUT_JSON.write_text(
        json.dumps(
            {
                "meta": report,
                "model": {"proximity_levels": list(PROXIMITY_LEVELS)},
                "records": records,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    OUT_REPORT.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    refresh_date = report["generated_at"][:10]
    intro = load_intro()
    OUT_MD.write_text(
        f'''---
hide:
  - toc
---
# Revues {{ .page-title-compact }}

<div class="revues-intro" markdown="1">

{intro}

</div>

<p class="revues-source">Grappe Mir@bel n° {GRAPPE_ID} « {GRAPPE_NAME} ».</p>
<p class="revues-source-link"><a href="{GRAPPE_URL}" target="_blank" rel="noopener">Consulter la grappe sur Mir@bel</a></p>
<p class="revues-updated">Dernière actualisation : <strong>{refresh_date}</strong> (API Mir@bel).</p>

<link rel="stylesheet" href="../assets/stylesheets/revues-exploration.css">

<div id="revues-explorer" class="revues-explorer" data-source="../assets/data/revues-exploration.json"><p class="explorer-loading">Chargement des revues…</p></div>

<script src="../assets/javascripts/revues-config.js" defer></script>
<script src="../assets/javascripts/revues-exploration.js" defer></script>
''',
        encoding="utf-8",
    )
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__ == "__main__": main()