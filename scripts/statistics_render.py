import csv
import html
import math

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from statistics_data import rpct
from statistics_publication import downloads_for

BLUES = ["#07558c", "#367fb8", "#67a9df", "#a7d5f5", "#8fc1e8", "#bdddf4", "#4d91c7", "#245f92"]
OPEN = {"open": "#07558c", "restricted": "#a7d5f5"}


def esc(value):
    return html.escape(str(value), quote=True)


def fmt_pct(value, total, lang):
    value = rpct(value, total)
    return f"{value} %" if lang == "fr" else f"{value}%"


def icon(ext):
    """Inline decorative SVG; the format remains accessible HTML text below."""
    shapes = {
        "csv": '<path d="M8 3h19l11 11v31H8z"/><path d="M27 3v11h11"/><path d="M15 23h16M15 29h16M15 35h11"/>',
        "xlsx": '<rect x="8" y="5" width="30" height="39" rx="2"/><path d="M8 17h30M8 27h30M8 37h30M23 17v27"/>',
        "jpg": '<rect x="6" y="7" width="34" height="34" rx="3"/><circle cx="17" cy="17" r="3"/><path d="m9 35 11-11 7 7 5-5 6 9"/>',
        "svg": '<path d="M8 3h19l11 11v31H8z"/><path d="M27 3v11h11"/><circle cx="16" cy="33" r="2"/><circle cx="30" cy="23" r="2"/><circle cx="30" cy="37" r="2"/><path d="M18 32c5-11 10-10 10-9M18 34c5 7 9 5 10 3"/>',
    }
    return (
        '<span class="chart-download-icon" aria-hidden="true">'
        '<svg viewBox="0 0 46 48" focusable="false" fill="none" '
        'stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round">'
        + shapes[ext] + '</svg></span>'
    )


def actions(chart_id, lang):
    base = "../downloads/" if lang == "fr" else "../../downloads/en/"
    groups = (
        (("Données", ("csv", "xlsx")), ("Graphique", ("jpg", "svg")))
        if lang == "fr" else
        (("Data", ("csv", "xlsx")), ("Chart", ("jpg", "svg")))
    )

    def group(label, formats):
        links = []
        for ext in formats:
            description = (
                ("Télécharger les données" if ext in ("csv", "xlsx") else "Télécharger le graphique")
                if lang == "fr" else
                ("Download data" if ext in ("csv", "xlsx") else "Download chart")
            )
            readable = ext.upper()
            links.append(
                f'<a class="chart-download" href="{base}{chart_id}.{ext}" download '
                f'aria-label="{description} ({readable})" title="{description} ({readable})">'
                f'{icon(ext)}<span class="chart-download-format">{readable}</span></a>'
            )
        return (
            '<div class="chart-actions-row">'
            f'<span class="chart-actions-label">{label} :</span>'
            '<div class="chart-actions-formats">' + "".join(links) + '</div>'
            '</div>'
        )

    return '<div class="chart-actions">' + "".join(group(label, formats) for label, formats in groups) + '</div>'


def pie_legend(labels, values, lang, sort_by_count=True):
    """HTML legend, sorted visually without altering sector positions or colors."""
    total = sum(values)
    items = list(enumerate(zip(labels, values)))
    if sort_by_count:
        items.sort(key=lambda item: -item[1][1])  # stable on ties
    entries = []
    for i, (label, value) in items:
        entries.append(
            f'<span class="chart-legend-item"><span class="chart-legend-swatch" '
            f'style="background:{BLUES[i % len(BLUES)]}" aria-hidden="true"></span>'
            f'<span>{esc(label)} — {value} ({fmt_pct(value, total, lang)})</span></span>'
        )
    return '<div class="chart-legend" role="group" aria-label="' + ('Légende' if lang == 'fr' else 'Legend') + '">' + ''.join(entries) + '</div>'


def stacked_legend(series, lang):
    entries = [
        f'<span class="chart-legend-item"><span class="chart-legend-swatch" '
        f'style="background:{color}" aria-hidden="true"></span><span>{esc(label)}</span></span>'
        for label, key, color in series
    ]
    return '<div class="chart-legend" role="group" aria-label="' + ('Légende' if lang == 'fr' else 'Legend') + '">' + ''.join(entries) + '</div>'


def pie_svg(labels, values, lang):
    total = sum(values)
    cx, cy, radius = 180, 180, 140
    angle = 0
    parts = []
    for i, (label, value) in enumerate(zip(labels, values)):
        sweep = value * 360 / total
        a1 = math.radians(angle - 90)
        a2 = math.radians(angle + sweep - 90)
        x1, y1 = cx + radius * math.cos(a1), cy + radius * math.sin(a1)
        x2, y2 = cx + radius * math.cos(a2), cy + radius * math.sin(a2)
        large = 1 if sweep > 180 else 0
        color = BLUES[i % len(BLUES)]
        parts.append(f'<path d="M {cx},{cy} L {x1:.2f},{y1:.2f} A {radius},{radius} 0 {large},1 {x2:.2f},{y2:.2f} Z" fill="{color}" stroke="#fff" stroke-width="2"/>')
        mid = math.radians(angle + sweep / 2 - 90)
        tx, ty = cx + radius * .62 * math.cos(mid), cy + radius * .62 * math.sin(mid)
        if sweep >= 18:
            parts.append(f'<text x="{tx:.1f}" y="{ty:.1f}" text-anchor="middle" style="font:700 14px Arial;fill:#08264a"><tspan x="{tx:.1f}">{value}</tspan><tspan x="{tx:.1f}" dy="18">{fmt_pct(value,total,lang)}</tspan></text>')
        angle += sweep

    return f'<svg style="display:block;width:100%;max-width:540px;height:auto" viewBox="0 0 360 360">' + "".join(parts) + "</svg>"


def stacked_svg(categories, series, values, lang, percent=False, average=None):
    width, height, left, right, top, bottom = 900, 520, 70, 30, 30, 145
    plot_width = width - left - right
    plot_height = height - top - bottom
    gap = plot_width / max(len(categories), 1)
    bar_width = min(115, gap * .62)
    max_total = max(sum(values[c].values()) for c in categories) if categories else 1
    parts = []

    for i, category in enumerate(categories):
        x = left + i * gap + (gap - bar_width) / 2
        base = top + plot_height
        total = sum(values[category].values())
        for label, key, color in series:
            value = values[category].get(key, 0)
            height_value = plot_height * (value / total if percent and total else value / max_total)
            y = base - height_value
            parts.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{bar_width:.1f}" height="{height_value:.1f}" fill="{color}"/>')
            text = fmt_pct(value, total, lang) if percent else str(value)
            if height_value > 25:
                text_color = "#fff" if key == "open" else "#08264a"
                parts.append(f'<text x="{x+bar_width/2:.1f}" y="{y+height_value/2+5:.1f}" text-anchor="middle" style="font:700 14px Arial;fill:{text_color}">{text}</text>')
            base = y
        tx = x + bar_width / 2
        parts.append(f'<text x="{tx:.1f}" y="{top+plot_height+22}" text-anchor="end" transform="rotate(-32 {tx:.1f},{top+plot_height+22})" style="font:13px Arial;fill:#08264a">{esc(category)}</text>')

    parts.append(f'<line x1="{left}" y1="{top+plot_height}" x2="{width-right}" y2="{top+plot_height}" stroke="#dce7f2"/>')
    if average is not None and percent:
        ay = top + plot_height * (1 - average)
        opened = round(average * 100)
        restricted = 100 - opened
        label = (f"Moyenne sur l’ensemble des revues du périmètre ({opened} % en accès ouvert, {restricted} % en accès restreint)" if lang == "fr" else f"Average across all journals in the corpus ({opened}% open access, {restricted}% restricted access)")
        parts.append(f'<line x1="{left}" y1="{ay:.1f}" x2="{width-right}" y2="{ay:.1f}" stroke="#b74f17" stroke-width="2" stroke-dasharray="7 5"/><text x="{left+5}" y="{max(top+14,ay-7):.1f}" style="font:13px Arial;fill:#08264a">{esc(label)}</text>')

    return f'<svg style="display:block;width:100%;max-width:900px;height:auto" viewBox="0 0 {width} {height}">' + "".join(parts) + "</svg>"


def bar_svg(labels, values):
    width, height, left, right, top, bottom = 960, 430, 55, 25, 25, 95
    plot_width = width - left - right
    plot_height = height - top - bottom
    max_value = max(values) if values else 1
    gap = plot_width / max(len(labels), 1)
    bar_width = max(5, gap * .68)
    parts = []
    for i, (label, value) in enumerate(zip(labels, values)):
        x = left + i * gap + (gap - bar_width) / 2
        bar_height = plot_height * value / max_value
        y = top + plot_height - bar_height
        parts.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{bar_width:.1f}" height="{bar_height:.1f}" fill="{BLUES[i % 4]}"/><text x="{x+bar_width/2:.1f}" y="{max(y-5,12):.1f}" text-anchor="middle" style="font:700 12px Arial;fill:#08264a">{value}</text>')
        if len(labels) <= 15 or i % max(1, math.ceil(len(labels) / 12)) == 0:
            parts.append(f'<text x="{x+bar_width/2:.1f}" y="{top+plot_height+18}" text-anchor="end" transform="rotate(-45 {x+bar_width/2:.1f},{top+plot_height+18})" style="font:12px Arial;fill:#08264a">{esc(label)}</text>')
    return f'<svg style="display:block;width:100%;max-width:960px;height:auto" viewBox="0 0 {width} {height}">' + "".join(parts) + "</svg>"


def export_csv(chart_id, headers, rows, lang="fr"):
    """Export the same tabular data as CSV and as a typed Excel workbook."""
    destination = downloads_for(lang)
    destination.mkdir(parents=True, exist_ok=True)
    rows = list(rows)

    with (destination / f"{chart_id}.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(headers)
        writer.writerows(rows)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Données" if lang == "fr" else "Data"
    sheet.append(list(headers))
    for row in rows:
        sheet.append([_excel_value(value) for value in row])

    header_fill = PatternFill(fill_type="solid", fgColor="08264A")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        width = min(60, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[letter].width = width
    workbook.save(destination / f"{chart_id}.xlsx")


def _excel_value(value):
    """Keep genuine numbers numeric; leave textual labels unchanged."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    return value


def export_svg(chart_id, svg, lang="fr"):
    destination = downloads_for(lang)
    destination.mkdir(parents=True, exist_ok=True)
    (destination / f"{chart_id}.svg").write_text('<?xml version="1.0" encoding="UTF-8"?>\n' + svg.replace("<svg ", '<svg xmlns="http://www.w3.org/2000/svg" ', 1), encoding="utf-8")


def export_jpg_pie(chart_id, labels, values, lang="fr"):
    destination = downloads_for(lang)
    destination.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(8, 5), dpi=150)
    ax.pie(values, labels=labels, autopct=lambda p: f"{round(p)}%", colors=[BLUES[i % len(BLUES)] for i in range(len(values))], startangle=90)
    ax.axis("equal")
    fig.savefig(destination / f"{chart_id}.jpg", bbox_inches="tight", facecolor="white")
    plt.close(fig)


def export_jpg_bar(chart_id, labels, values, lang="fr"):
    destination = downloads_for(lang)
    destination.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(10, 5), dpi=150)
    ax.bar(range(len(values)), values)
    ax.set_xticks(range(len(labels)), labels, rotation=45, ha="right")
    fig.tight_layout()
    fig.savefig(destination / f"{chart_id}.jpg", facecolor="white")
    plt.close(fig)


def export_jpg_stacked(chart_id, categories, series, values, percent=False, average=None, lang="fr"):
    destination = downloads_for(lang)
    destination.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(9, 5), dpi=150)
    bottom = [0] * len(categories)
    for label, key, color in series:
        ys = []
        for category in categories:
            total = sum(values[category].values())
            ys.append(values[category].get(key, 0) * 100 / total if percent and total else values[category].get(key, 0))
        ax.bar(categories, ys, bottom=bottom, label=label, color=color)
        bottom = [a + b for a, b in zip(bottom, ys)]
    if average is not None and percent:
        ax.axhline(average * 100, linestyle="--")
    ax.legend()
    ax.tick_params(axis="x", rotation=30)
    fig.tight_layout()
    fig.savefig(destination / f"{chart_id}.jpg", facecolor="white")
    plt.close(fig)
