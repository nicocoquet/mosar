"""Acquisition des sources de données utilisées par l'explorateur des revues."""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import Request, urlopen


@dataclass(frozen=True)
class MirabelSource:
    """Configuration d'accès aux données Mir@bel."""

    api: str
    cluster_id: int
    project_name: str
    site_url: str


@dataclass(frozen=True)
class MosarJournalSource:
    """Configuration de la source analytique Mosar."""

    xlsx: Path
    sheet: str
    columns: dict[str, str]
    project_name: str
    fallback_enabled: bool
    fallback: Path


def _api_json(source: MirabelSource, path: str):
    """Interroger un endpoint JSON de l'API Mir@bel."""

    request = Request(
        f"{source.api.rstrip('/')}{path}",
        headers={
            "User-Agent": (
                f"{source.project_name}/1.0 (+{source.site_url})"
            )
        },
    )

    with urlopen(request, timeout=45) as response:
        return json.load(response)


def fetch_mirabel_titles(source: MirabelSource) -> list[dict]:
    """Charger les titres actifs de la grappe Mir@bel."""

    records: list[dict] = []
    offset = 0

    while True:
        query = urlencode(
            {
                "grappeid": source.cluster_id,
                "actif": 1,
                "offset": offset,
            }
        )

        batch = _api_json(source, f"/titres?{query}")

        if not isinstance(batch, list):
            raise RuntimeError("Réponse inattendue de /titres")

        records.extend(batch)

        if len(batch) < 1000:
            return records

        offset += 1000


def fetch_mirabel_themes(source: MirabelSource) -> dict[int, list[str]]:
    """Charger les thèmes Mir@bel et les indexer par identifiant de revue."""

    payload = _api_json(
        source,
        f"/themes/grappe/{source.cluster_id}",
    )

    if not isinstance(payload, list):
        raise RuntimeError("Réponse inattendue de /themes/grappe")

    return {
        int(item["revueid"]): [
            theme.get("nom")
            for theme in item.get("themes", [])
            if theme.get("nom")
        ]
        for item in payload
    }


def load_mosar_journals(
    source: MosarJournalSource,
) -> tuple[dict[int, dict], str]:
    """Charger les enrichissements Mosar depuis le XLSX ou son fallback CSV."""

    try:
        from openpyxl import load_workbook

        worksheet = load_workbook(
            source.xlsx,
            data_only=True,
            read_only=True,
        )[source.sheet]

        rows = worksheet.iter_rows(values_only=True)

        headers = [
            str(value).strip() if value is not None else ""
            for value in next(rows)
        ]

        mirabel_column = source.columns["mirabel_id"]

        if mirabel_column in headers:
            positions = {
                header: index
                for index, header in enumerate(headers)
            }

            fields = {
                "nom_revue": source.columns["journal_name"],
                "sous_titre": source.columns["journal_subtitle"],
                "niveau_rattachement": source.columns["proximity"],
                "acces_ouvert": source.columns["open_access"],
                "discipline": source.columns["discipline"],
                "periodicite": source.columns["periodicity"],
                "format": source.columns["format"],
                "type_editeur": source.columns["publisher_type"],
                "licence": source.columns["licence"],
            }

            output: dict[int, dict] = {}

            for row in rows:
                raw_id = row[positions[mirabel_column]]

                if raw_id in (None, ""):
                    continue

                output[int(raw_id)] = {
                    key: (
                        row[positions[column]]
                        if column in positions
                        else None
                    )
                    for key, column in fields.items()
                }

            return output, f"XLSX {source.project_name}"

    except Exception as exc:
        if not source.fallback_enabled:
            raise RuntimeError(
                "Lecture XLSX impossible pour la jointure Mosar "
                "et fallback CSV désactivé"
            ) from exc

        print(
            f"Lecture XLSX impossible pour la jointure ({exc}); "
            f"recours au CSV configuré : {source.fallback}."
        )

    output = {}

    with source.fallback.open(
        encoding="utf-8",
        newline="",
    ) as handle:
        for row in csv.DictReader(handle):
            mirabel_id = int(row.pop("mirabel_id"))
            output[mirabel_id] = row

    return output, f"CSV de fallback ({source.fallback})"