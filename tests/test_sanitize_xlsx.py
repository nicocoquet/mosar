"""Tests autonomes de l'expurgation XLSX : python -m unittest discover -s tests -p 'test_sanitize_xlsx.py'."""
import tempfile
import unittest
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from sanitize_xlsx import sanitize, validate_sanitized


class SanitizeXlsxTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.source = self.root / "source.xlsx"
        self.output = self.root / "output.xlsx"

    def make_workbook(self, headers, rows=(), extra_sheets=()):
        wb = Workbook()
        ws = wb.active
        ws.title = "recensement"
        ws.append(headers)
        for row in rows:
            ws.append(row)
        for name in extra_sheets:
            wb.create_sheet(name).append(["information à exclure"])
        wb.save(self.source)
        wb.close()

    def test_removes_other_sheets_private_columns_and_empty_ids(self):
        self.make_workbook(
            ["ID_Mosar", "Nom", "Contacts", "Divers"],
            [[1, "Revue A", "privé", "privé"], [None, "Brouillon", "privé", "privé"]],
            ["Notes", "Archive"],
        )
        sanitize(self.source, self.output)
        wb = load_workbook(self.output, read_only=True, data_only=True)
        self.assertEqual(wb.sheetnames, ["recensement"])
        self.assertEqual(list(wb.active.values), [("ID_Mosar", "Nom"), (1, "Revue A")])
        wb.close()

    def test_accepts_already_sanitized_workbook(self):
        self.make_workbook(["ID_Mosar", "Nom"], [[2, "Revue B"]])
        sanitize(self.source, self.output)
        validate_sanitized(self.output)

    def test_missing_id_column_is_explicit(self):
        self.make_workbook(["Nom"], [["Revue A"]])
        with self.assertRaisesRegex(SystemExit, "Colonne obligatoire introuvable"):
            sanitize(self.source, self.output)
        self.assertFalse(self.output.exists())

    def test_missing_sheet_is_explicit(self):
        wb = Workbook()
        wb.active.title = "autre"
        wb.save(self.source)
        wb.close()
        with self.assertRaisesRegex(SystemExit, "Onglet attendu introuvable"):
            sanitize(self.source, self.output)

    def test_check_rejects_extra_sheet(self):
        self.make_workbook(["ID_Mosar"], [[1]], ["secret"])
        with self.assertRaisesRegex(SystemExit, "uniquement l’onglet"):
            validate_sanitized(self.source)

    def test_check_rejects_private_column(self):
        self.make_workbook(["ID_Mosar", "Contacts"], [[1, "secret"]])
        with self.assertRaisesRegex(SystemExit, "Colonnes interdites"):
            validate_sanitized(self.source)

    def test_check_rejects_empty_id(self):
        self.make_workbook(["ID_Mosar", "Nom"], [[None, "Revue"]])
        with self.assertRaisesRegex(SystemExit, "ID_Mosar vide"):
            validate_sanitized(self.source)

    def test_no_valid_rows_does_not_replace_existing_output(self):
        self.make_workbook(["ID_Mosar", "Nom"], [[None, "Brouillon"]])
        self.output.write_bytes(b"existing")
        with self.assertRaisesRegex(SystemExit, "Aucune revue"):
            sanitize(self.source, self.output)
        self.assertEqual(self.output.read_bytes(), b"existing")

    def test_rejects_in_place_rewrite(self):
        self.make_workbook(["ID_Mosar"], [[1]])
        with self.assertRaisesRegex(SystemExit, "doivent être différents"):
            sanitize(self.source, self.source)


if __name__ == "__main__":
    unittest.main()
