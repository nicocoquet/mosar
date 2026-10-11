"""Tests API Mosar, sans OAuth externe ni écriture GitHub réelle."""
import base64
import io
import sys
import unittest
from pathlib import Path
from urllib.parse import urlparse, parse_qs
from unittest.mock import AsyncMock, patch

from fastapi.testclient import TestClient
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
CONFIG = {"data": {"sheet": "recensement", "file": "data/Recensement-revues-Mosar.xlsx"},
          "privacy": {"allowed_columns": ["ID_Mosar", "Nom"], "forbidden_columns": ["Contacts"]}}


def make_xlsx(name="Revue", secret="secret"):
    wb = Workbook()
    ws = wb.active
    ws.title = "recensement"
    ws.append(["ID_Mosar", "Nom", "Contacts"])
    ws.append(["M-1", name, secret])
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


class ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Importer la vraie configuration et ses fonctions utilitaires : les
        # modules statistiques en dépendent dès l'import de service.app.
        sys.path.insert(0, str(ROOT / "scripts"))
        from service.app import Settings, create_app
        cls.settings_class = Settings
        cls.create_app = staticmethod(create_app)

    def setUp(self):
        settings = self.settings_class("id", "secret", "x" * 48)
        # Les tests historiques utilisent volontairement un classeur minimal.
        # On isole ici le moteur statistique, testé séparément, sans neutraliser
        # l'assainissement, l'authentification ni la logique de publication.
        self.config_patch = patch("service.app.load_config", return_value=CONFIG)
        self.config_patch.start()
        self.addCleanup(self.config_patch.stop)
        self.rows_patch = patch("service.app.read_rows", return_value=[])
        self.rows_mock = self.rows_patch.start()
        self.addCleanup(self.rows_patch.stop)
        self.derive_patch = patch("service.app.derive", return_value={})
        self.derive_mock = self.derive_patch.start()
        self.addCleanup(self.derive_patch.stop)
        self.build_patch = patch("service.app.build", return_value="Contenu valide")
        self.build_mock = self.build_patch.start()
        self.addCleanup(self.build_patch.stop)

        self.app = self.create_app(settings)
        self.client = TestClient(self.app)
        self.addCleanup(self.client.close)
        self.key = self.app.state.sessions.create("nicocoquet", "fake-github-token")
        self.headers = {"Authorization": f"Bearer {self.key}"}

    def payload(self, name="Revue", secret="secret"):
        return {"filename": "recensement.xlsx", "content": base64.b64encode(make_xlsx(name, secret)).decode()}

    def oauth_state(self, destination):
        result = self.client.get("/auth/github", params={"returnTo": destination}, follow_redirects=False)
        self.assertEqual(result.status_code, 307)
        return parse_qs(urlparse(result.headers["location"]).query)["state"][0]

    def test_oauth_destinations(self):
        for destination in ("https://nicocoquet.github.io/mosar/admin/", "https://nicocoquet.github.io/mosar/en/admin/"):
            with self.subTest(destination=destination):
                self.assertTrue(self.oauth_state(destination))
        for destination in ("https://evil.example/mosar/admin/", "https://nicocoquet.github.io/mosar/", "https://nicocoquet.github.io/mosar/admin/?next=evil", "https://nicocoquet.github.io/mosar/admin/#fragment", "http://nicocoquet.github.io/mosar/admin/"):
            with self.subTest(destination=destination):
                response = self.client.get("/auth/github", params={"returnTo": destination}, follow_redirects=False)
                self.assertEqual(response.status_code, 400)

    def test_oauth_callback_one_use_state(self):
        from unittest.mock import MagicMock
        destination = "https://nicocoquet.github.io/mosar/en/admin/"
        state = self.oauth_state(destination)
        fake_response = MagicMock()
        fake_response.json.return_value = {"access_token": "fake-token"}
        fake_response.raise_for_status.return_value = None
        fake_client = AsyncMock()
        fake_client.post.return_value = fake_response
        fake_context = AsyncMock()
        fake_context.__aenter__.return_value = fake_client
        with patch("service.app.httpx.AsyncClient", return_value=fake_context), patch("service.app.github_request", new_callable=AsyncMock, return_value={"login": "nicocoquet"}):
            response = self.client.get("/auth/callback", params={"code": "code", "state": state}, follow_redirects=False)
            self.assertEqual(response.status_code, 303, response.text)
            self.assertTrue(response.headers["location"].startswith(destination + "#session="))
            replay = self.client.get("/auth/callback", params={"code": "code", "state": state}, follow_redirects=False)
            self.assertEqual(replay.status_code, 400)

    def test_logout_revokes_session(self):
        self.assertEqual(self.client.get("/auth/session", headers=self.headers).status_code, 200)
        self.assertEqual(self.client.post("/auth/logout", headers=self.headers).status_code, 200)
        self.assertEqual(self.client.get("/auth/session", headers=self.headers).status_code, 401)

    def test_health(self):
        self.assertEqual(self.client.get("/health").json(), {"status": "ok"})

    def test_auth_required(self):
        self.assertEqual(self.client.post("/api/preview", json=self.payload()).status_code, 401)
        self.assertEqual(self.client.post("/api/publish", json={**self.payload(), "preview_token": "bad"}).status_code, 401)

    def test_preview_removes_secrets_from_report(self):
        response = self.client.post("/api/preview", json=self.payload(secret="private-contact-value"), headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["report"]["private_columns_removed"], ["Contacts"])
        self.assertNotIn("private-contact-value", response.text)
        self.assertEqual(response.json()["report"]["rows_retained"], 1)

    def test_preview_checks_sanitized_statistics_in_both_languages(self):
        response = self.client.post("/api/preview", json=self.payload(), headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        self.rows_mock.assert_called_once()
        source = self.rows_mock.call_args.args[0]
        self.assertEqual(source.xlsx.name, "recensement.xlsx")
        self.derive_mock.assert_called_once()
        self.assertEqual(self.build_mock.call_count, 2)
        self.assertEqual([call.args[1] for call in self.build_mock.call_args_list], ["fr", "en"])

    def test_invalid_statistics_refused_before_github(self):
        self.derive_mock.side_effect = SystemExit("Périodicité invalide ligne 93: 'inconnue'")
        with patch("service.app.github_request", new_callable=AsyncMock) as github:
            response = self.client.post("/api/preview", json=self.payload(), headers=self.headers)
            self.assertEqual(response.status_code, 422, response.text)
            self.assertIn("Périodicité invalide ligne 93", response.json()["detail"])
            github.assert_not_awaited()

    def test_invalid_page_build_refused_before_github(self):
        self.build_mock.side_effect = ValueError("Statistiques non générables")
        with patch("service.app.github_request", new_callable=AsyncMock) as github:
            response = self.client.post("/api/preview", json=self.payload(), headers=self.headers)
            self.assertEqual(response.status_code, 422, response.text)
            github.assert_not_awaited()

    def test_publish_rechecks_statistics_before_github(self):
        response = self.client.post("/api/preview", json=self.payload(), headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        token = response.json()["preview_token"]
        self.derive_mock.side_effect = SystemExit("Année de création invalide ligne 2")
        with patch("service.app.github_request", new_callable=AsyncMock) as github:
            result = self.client.post(
                "/api/publish", json={**self.payload(), "preview_token": token},
                headers=self.headers,
            )
            self.assertEqual(result.status_code, 422, result.text)
            github.assert_not_awaited()

    def test_invalid_xlsx_rejected(self):
        response = self.client.post("/api/preview", json={"filename": "file.xlsx", "content": base64.b64encode(b"bad").decode()}, headers=self.headers)
        self.assertEqual(response.status_code, 422)

    def test_invalid_name_rejected(self):
        response = self.client.post("/api/preview", json={**self.payload(), "filename": "../file.xlsx"}, headers=self.headers)
        self.assertEqual(response.status_code, 400)

    def test_changed_file_cannot_be_published(self):
        response = self.client.post("/api/preview", json=self.payload(), headers=self.headers)
        token = response.json()["preview_token"]
        with patch("service.app.github_request", new_callable=AsyncMock) as github:
            result = self.client.post("/api/publish", json={**self.payload(name="Autre revue"), "preview_token": token}, headers=self.headers)
            self.assertEqual(result.status_code, 409, result.text)
            github.assert_not_awaited()

    def test_publish_only_sanitized_content(self):
        from openpyxl import load_workbook
        response = self.client.post("/api/preview", json=self.payload(), headers=self.headers)
        token = response.json()["preview_token"]
        async def fake_github(method, path, access_token, **kwargs):
            if method == "GET":
                return {"sha": "oldsha"}
            self.assertEqual(kwargs["json"]["sha"], "oldsha")
            content = base64.b64decode(kwargs["json"]["content"])
            wb = load_workbook(io.BytesIO(content), read_only=True)
            try:
                self.assertEqual([c.value for c in wb.active[1]], ["ID_Mosar", "Nom"])
                self.assertEqual(wb.active.max_row, 2)
            finally:
                wb.close()
            return {"commit": {"sha": "newsha"}}
        with patch("service.app.github_request", side_effect=fake_github):
            result = self.client.post("/api/publish", json={**self.payload(), "preview_token": token}, headers=self.headers)
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()["commit_sha"], "newsha")

    def test_other_user_cannot_publish(self):
        response = self.client.post("/api/preview", json=self.payload(), headers=self.headers)
        other = self.app.state.sessions.create("colleague", "other-token")
        result = self.client.post("/api/publish", json={**self.payload(), "preview_token": response.json()["preview_token"]}, headers={"Authorization": f"Bearer {other}"})
        self.assertEqual(result.status_code, 403)

    def test_job_status(self):
        response = self.client.post("/api/preview", json=self.payload(), headers=self.headers)
        with patch("service.app.github_request", new_callable=AsyncMock) as github:
            github.side_effect = [{"sha": "old"}, {"commit": {"sha": "new"}}]
            published = self.client.post("/api/publish", json={**self.payload(), "preview_token": response.json()["preview_token"]}, headers=self.headers)
            self.assertEqual(published.status_code, 200)
            github.reset_mock()
            github.side_effect = None
            github.return_value = {"workflow_runs": [{"status": "completed", "conclusion": "success", "html_url": "https://github.com/example"}]}
            status = self.client.get("/api/jobs/" + published.json()["job_token"], headers=self.headers)
            self.assertEqual(status.json()["status"], "completed")


if __name__ == "__main__":
    unittest.main()
