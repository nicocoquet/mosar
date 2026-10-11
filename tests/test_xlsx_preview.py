"""Tests unitaires autonomes du moteur de prévisualisation Mosar."""
import io
import sys
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from xlsx_preview import prepare_xlsx

ALLOWED = {'ID_Mosar', 'Nom', 'Site web'}
FORBIDDEN = {'Contacts', 'Divers'}


def source(headers, rows, extra=False):
    wb = Workbook()
    ws = wb.active
    ws.title = 'recensement'
    ws.append(headers)
    for row in rows:
        ws.append(row)
    if extra:
        wb.create_sheet('Privé').append(['confidentiel'])
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def prepare(content):
    return prepare_xlsx(content, sheet='recensement', allowed_columns=ALLOWED,
                        forbidden_columns=FORBIDDEN)


class PreviewTests(unittest.TestCase):
    def test_private_unknown_and_sheets_removed(self):
        result = prepare(source(['ID_Mosar', 'Nom', 'Contacts', 'Inconnue'],
                                [['M1', 'Revue', 'secret', 'autre secret']], True))
        self.assertEqual(result.report['private_columns_removed'], ['Contacts'])
        self.assertEqual(result.report['unknown_columns_removed'], ['Inconnue'])
        self.assertEqual(result.report['sheets_removed'], ['Privé'])
        self.assertNotIn(b'secret', result.content)
        wb = load_workbook(io.BytesIO(result.content))
        self.assertEqual(wb.sheetnames, ['recensement'])
        self.assertEqual([c.value for c in wb.active[1]], ['ID_Mosar', 'Nom'])
        wb.close()

    def test_formula_in_published_column_refused(self):
        with self.assertRaisesRegex(ValueError, 'Formule Excel interdite'):
            prepare(source(['ID_Mosar', 'Nom'], [['M1', '=1+1']]))

    def test_formula_in_removed_column_ignored(self):
        result = prepare(source(['ID_Mosar', 'Nom', 'Contacts'], [['M1', 'Revue', '=1+1']]))
        self.assertEqual(result.report['rows_retained'], 1)

    def test_rows_without_id_removed(self):
        result = prepare(source(['ID_Mosar', 'Nom'], [['M1', 'Revue'], ['', 'Sans ID']]))
        self.assertEqual(result.report['rows_without_id_removed'], 1)
        self.assertEqual(result.report['rows_retained'], 1)

    def test_missing_id_rejected(self):
        with self.assertRaisesRegex(ValueError, 'ID_Mosar'):
            prepare(source(['Nom'], [['Revue']]))

    def test_duplicate_headers_rejected(self):
        with self.assertRaisesRegex(ValueError, 'dupliqués'):
            prepare(source(['ID_Mosar', 'Nom', 'Nom'], [['M1', 'a', 'b']]))

    def test_no_retained_rows_rejected(self):
        with self.assertRaisesRegex(ValueError, 'Aucune revue'):
            prepare(source(['ID_Mosar', 'Nom'], [['', 'a']]))

    def test_non_xlsx_rejected(self):
        with self.assertRaisesRegex(ValueError, 'Impossible de lire'):
            prepare(b'not an xlsx')

    def test_no_private_values_in_report(self):
        result = prepare(source(['ID_Mosar', 'Nom', 'Divers'], [['M1', 'Revue', 'SECRET-NEVER-REPORT']]))
        self.assertNotIn('SECRET-NEVER-REPORT', repr(result.report))


if __name__ == '__main__':
    unittest.main()
